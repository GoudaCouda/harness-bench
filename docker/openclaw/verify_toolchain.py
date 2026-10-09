#!/usr/bin/env python3
"""Fail closed unless the shared task-authoring toolchain is usable."""

from __future__ import annotations

import argparse
import importlib.metadata
import json
import sqlite3
import subprocess
import tempfile
from pathlib import Path

import docx
import openpyxl
import pptx
import pypdf
import pytest
import slugify
import yaml


EXPECTED = {
    "openpyxl": "3.1.5",
    "pypdf": "5.3.1",
    "pytest": "8.3.5",
    "python-docx": "1.1.2",
    "python-pptx": "1.0.2",
    "python-slugify": "8.0.4",
    "PyYAML": "6.0.2",
}


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--python-only",
        action="store_true",
        help="verify the shared Python task tools without requiring amtool",
    )
    args = parser.parse_args()
    actual = {name: importlib.metadata.version(name) for name in EXPECTED}
    if actual != EXPECTED:
        raise SystemExit(
            "task toolchain version mismatch: "
            + json.dumps({"expected": EXPECTED, "actual": actual}, sort_keys=True)
        )

    with tempfile.TemporaryDirectory() as temporary:
        root = Path(temporary)

        presentation = pptx.Presentation()
        presentation.slides.add_slide(presentation.slide_layouts[0])
        presentation.save(root / "probe.pptx")
        if len(pptx.Presentation(root / "probe.pptx").slides) != 1:
            raise SystemExit("python-pptx round trip failed")

        document = docx.Document()
        document.add_paragraph("HarnessBench")
        document.save(root / "probe.docx")
        if not docx.Document(root / "probe.docx").paragraphs:
            raise SystemExit("python-docx round trip failed")

        workbook = openpyxl.Workbook()
        workbook.active["A1"] = "HarnessBench"
        workbook.save(root / "probe.xlsx")
        if (
            openpyxl.load_workbook(root / "probe.xlsx").active["A1"].value
            != "HarnessBench"
        ):
            raise SystemExit("openpyxl round trip failed")

        writer = pypdf.PdfWriter()
        writer.add_blank_page(width=72, height=72)
        with (root / "probe.pdf").open("wb") as stream:
            writer.write(stream)
        if len(pypdf.PdfReader(root / "probe.pdf").pages) != 1:
            raise SystemExit("pypdf round trip failed")

    if slugify.slugify("Harness Bench") != "harness-bench":
        raise SystemExit("python-slugify probe failed")
    if yaml.safe_load("enabled: true") != {"enabled": True}:
        raise SystemExit("PyYAML probe failed")
    if not pytest.__version__.startswith("8.3."):
        raise SystemExit("pytest probe failed")
    with sqlite3.connect(":memory:") as connection:
        if connection.execute("SELECT 1").fetchone() != (1,):
            raise SystemExit("sqlite3 in-memory query failed")
    sqlite_cli = subprocess.run(
        ["sqlite3", ":memory:", "SELECT 1;"],
        check=False,
        capture_output=True,
        text=True,
        timeout=10,
    )
    if sqlite_cli.returncode != 0 or sqlite_cli.stdout.strip() != "1":
        raise SystemExit("sqlite3 CLI probe failed")

    if not args.python_only:
        amtool = subprocess.run(
            ["amtool", "--version"],
            check=False,
            capture_output=True,
            text=True,
            timeout=10,
        )
        amtool_version = (amtool.stdout + amtool.stderr).strip()
        if amtool.returncode != 0 or "version 0.32.1" not in amtool_version:
            raise SystemExit("amtool 0.32.1 probe failed")

    print(
        json.dumps(
            {
                "status": "pass",
                "packages": actual,
                "sqlite3_cli": "available",
                "amtool": None if args.python_only else "0.32.1",
            },
            sort_keys=True,
        )
    )


if __name__ == "__main__":
    main()
